from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
OUTPUT_PATH = os.path.join(BASE, 'rapport_professionnel.docx')

code_examples = {
    'routes_upload': [
        'if fichier and extension_autorisee(fichier.filename):',
        '    nom_securise = secure_filename(fichier.filename)',
        '    chemin = os.path.join(current_app.config["UPLOAD_FOLDER"], nom_securise)',
        '    fichier.save(chemin)',
    ],
    'ocr_extraire_texte': [
        'if extension == "pdf":',
        '    pages = convert_from_path(chemin_fichier, dpi=300)',
        '    for page in pages:',
        '        texte_total += pytesseract.image_to_string(page, lang=langue) + "\n"',
    ],
    'model_document': [
        'class Document(db.Model):',
        '    nom_fichier = db.Column(db.String(255), nullable=False)',
        '    chemin_fichier = db.Column(db.String(500), nullable=False)',
        '    texte_ocr = db.Column(db.Text, nullable=True)',
        '    statut_ocr = db.Column(db.String(20), default="en_attente")',
    ],
}

sections = [
    {
        'title': 'Introduction',
        'content': [
            'Ce rapport professionnel présente le projet de site web de gestion des factures et documents techniques. Il est conçu comme une documentation d’ingénieur, structurée et illustrée par des extraits de code pour expliquer les choix techniques.',
            'Le rapport explique l’architecture globale, le workflow des imports, le traitement OCR, les fonctionnalités utilisateur et les points d’amélioration.',
        ],
        'placeholder': 'Capture d’écran page d’accueil / tableau de bord',
    },
    {
        'title': 'Architecture du projet',
        'content': [
            'Le projet utilise Flask comme framework web principal, avec Flask-SQLAlchemy pour la gestion de la base de données SQLite.',
            'Les composants frontend reposent sur Bootstrap 5 et une feuille de styles personnalisée pour assurer une expérience utilisateur moderne et responsive.',
            'La couche OCR s’appuie sur Tesseract via pytesseract, et les PDF sont convertis en images avec pdf2image pour permettre leur lecture.',
        ],
        'placeholder': 'Diagramme architecture technique',
    },
    {
        'title': 'Gestion des routes et du workflow',
        'content': [
            'La route /upload implémente l’upload de fichier, l’enregistrement sur le serveur et le traitement OCR immédiat.',
            'Le code vérifie l’extension du fichier, sauvegarde le document dans uploads, crée une entrée en base et lance le module OCR.',
            'L’approche garantit que l’utilisateur reçoit un retour rapide et que le document est persisté avant traitement.',
        ],
        'placeholder': 'Capture d’écran formulaire d’import',
        'code': code_examples['routes_upload'],
    },
    {
        'title': 'Module OCR et extraction de données',
        'content': [
            'Le module ocr_utils.py prend en charge les images et les fichiers PDF, en extrayant le texte brut puis en appliquant des règles d’analyse.',
            'L’extraction est basée sur des expressions régulières pour détecter les champs métier : numéro de facture, date, montant, fournisseur, type de panne et durée de maintenance.',
            'Cette solution est architecturée pour être réutilisable et évolutive.',
        ],
        'placeholder': 'Capture d’écran du flux OCR',
        'code': code_examples['ocr_extraire_texte'],
    },
    {
        'title': 'Modèle de données Document',
        'content': [
            'Le modèle Document structure les informations essentielles du document importé.',
            'Il inclut les métadonnées du fichier, le texte OCR, le statut de traitement et les champs métier extraits.',
            'Cette organisation facilite les recherches, l’affichage et l’export des documents.',
        ],
        'placeholder': 'Capture d’écran du modèle de données',
        'code': code_examples['model_document'],
    },
    {
        'title': 'Tableau de bord utilisateur',
        'content': [
            'Le tableau de bord fournit des statistiques clés et un tableau filtrable des documents importés.',
            'L’utilisateur peut trier, rechercher et filtrer par type, statut OCR et période.',
            'Les actions en masse permettent l’export CSV et la suppression multiple, améliorant la productivité.',
        ],
        'placeholder': 'Capture d’écran tableau de bord',
    },
    {
        'title': 'Page de détail document',
        'content': [
            'La page de détail affiche les informations extraites, le texte OCR brut et propose des champs modifiables.',
            'Cette page permet de corriger rapidement les erreurs d’extraction et de vérifier le document source.',
            'Une prévisualisation intégrée facilite le contrôle visuel du contenu sans téléchargement.',
        ],
        'placeholder': 'Capture d’écran page détail',
    },
    {
        'title': 'Assistant virtuel',
        'content': [
            'L’assistant virtuel est un composant interactif accessible depuis toutes les pages principales.',
            'Il guide les utilisateurs pour retrouver des documents ou comprendre des anomalies, avec un fallback local si OpenAI n’est pas configuré.',
            'Ce composant améliore l’ergonomie et réduit le temps de recherche dans l’application.',
        ],
        'placeholder': 'Capture d’écran assistant virtuel',
    },
    {
        'title': 'Sécurité et bonnes pratiques',
        'content': [
            'Le code inclut une vérification de session via @main_bp.before_request pour protéger les routes sensibles.',
            'Les mots de passe sont hachés et seules les pages de connexion restent publiques.',
            'Ces principes assurent une base solide pour la sécurité de l’application.',
        ],
        'placeholder': 'Capture d’écran de la page de connexion',
    },
    {
        'title': 'Installation et configuration technique',
        'content': [
            'Le projet s’installe avec Python 3, un environnement virtuel et les dépendances de requirements.txt.',
            'Tesseract et Poppler sont nécessaires pour l’OCR et la conversion des PDF.',
            'La configuration se fait via un fichier .env pour les paramètres sensibles et les chemins.',
        ],
        'placeholder': 'Capture d’écran configuration .env',
    },
    {
        'title': 'Tests et validation',
        'content': [
            'Le projet comprend des tests unitaires qui valident les routes et le comportement principal.',
            'L’utilisation de tests garantit que les modifications ne cassent pas le traitement de document.',
            'La validation est essentielle pour un projet d’ingénierie de qualité.',
        ],
        'placeholder': 'Capture d’écran des résultats de tests',
    },
    {
        'title': 'Recommandations d’évolution',
        'content': [
            'Pour un produit plus mature, il est recommandé d’ajouter la pagination, la gestion des rôles et l’archivage des documents.',
            'L’intégration de PDF.js améliorera l’expérience de prévisualisation des documents PDF.',
            'La migration vers PostgreSQL et un stockage objet est conseillée pour la production.',
        ],
        'placeholder': 'Illustration des évolutions',
    },
]


def add_title_page(doc):
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run('Rapport professionnel du projet Gestion Factures')
    run.bold = True
    run.font.size = Pt(28)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subrun = subtitle.add_run('Documentation technique avec extraits de code et zones de capture d’écran')
    subrun.italic = True
    subrun.font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph('Auteur : Ingénieur logiciel').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph('Date : 27 juillet 2026').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()


def add_section(doc, title, content, placeholder, code=None):
    doc.add_heading(title, level=1)
    for paragraph in content:
        p = doc.add_paragraph(paragraph)
        p.paragraph_format.space_after = Pt(8)
        p.runs[0].font.size = Pt(11)

    if code:
        doc.add_paragraph('Extrait de code :', style='Intense Quote')
        for line in code:
            code_line = doc.add_paragraph()
            run = code_line.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            code_line.paragraph_format.left_indent = Pt(12)
            code_line.paragraph_format.space_after = Pt(0)

    doc.add_paragraph()
    placeholder_paragraph = doc.add_paragraph(f'[Zone réservée pour capture d’écran : {placeholder}]')
    placeholder_paragraph.runs[0].italic = True
    placeholder_paragraph.runs[0].font.size = Pt(11)
    doc.add_page_break()


if __name__ == '__main__':
    document = Document()
    add_title_page(document)
    document.add_heading('Table des matières', level=1)
    for idx, section in enumerate(sections, start=1):
        line = document.add_paragraph(f'{idx}. {section["title"]}')
        line.paragraph_format.left_indent = Pt(12)
    document.add_page_break()

    for section in sections:
        add_section(document, section['title'], section['content'], section['placeholder'], section.get('code'))

    document.save(OUTPUT_PATH)
    print('Saved', OUTPUT_PATH)
