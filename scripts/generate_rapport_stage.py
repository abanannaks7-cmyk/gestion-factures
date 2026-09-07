from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
import os

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUTPUT_PATH = os.path.join(BASE, "rapport_stage_gestion_factures.docx")

sections = [
    {
        "title": "Introduction",
        "paragraphs": [
            "Dans le cadre de mon stage, j'ai participé à la conception et au développement d'une application web dédiée à la gestion automatisée des factures et des documents techniques. Ce projet répond à un besoin concret : centraliser les documents, réduire la saisie manuelle et faciliter la recherche d'informations importantes.",
            "L'application permet à un utilisateur authentifié d'importer un fichier PDF ou une image, d'en extraire le texte grâce à la reconnaissance optique de caractères (OCR), puis d'identifier automatiquement les principales informations de la facture.",
            "Ce rapport présente le contexte du projet, la démarche suivie, les choix techniques, les fonctionnalités réalisées, les difficultés rencontrées ainsi que les compétences acquises pendant le stage.",
        ],
        "screenshot": "Page de garde ou tableau de bord de l'application",
    },
    {
        "title": "Présentation de l'organisme d'accueil",
        "paragraphs": [
            "L'organisme d'accueil est présenté dans cette partie afin de situer le projet dans son environnement professionnel. Cette section doit être personnalisée avec le nom de la structure, son secteur d'activité, ses services principaux et son organisation.",
            "Le projet s'inscrit dans une démarche de dématérialisation et d'amélioration des processus administratifs. La gestion numérique des factures permet notamment de limiter les documents papier et de rendre les informations disponibles plus rapidement.",
            "La solution développée constitue un prototype fonctionnel pouvant être adapté aux règles et aux formats de documents utilisés par l'organisme.",
        ],
        "screenshot": "Logo ou présentation de l'organisme d'accueil",
    },
    {
        "title": "Contexte et problématique",
        "paragraphs": [
            "Le traitement manuel des factures demande du temps et peut entraîner des erreurs de saisie. Les documents sont parfois stockés dans plusieurs dossiers, ce qui rend leur consultation et leur suivi plus difficiles.",
            "La problématique du projet est donc la suivante : comment proposer une application simple permettant d'importer, d'analyser, de classer et de retrouver rapidement des factures et des documents techniques ?",
            "La réponse apportée consiste à associer une interface web, une base de données et un module OCR capable de transformer le contenu visuel d'un document en données exploitables.",
        ],
        "screenshot": "Schéma du problème avant et après la solution",
    },
    {
        "title": "Objectifs du stage",
        "paragraphs": [
            "L'objectif principal était de développer une application web fonctionnelle pour automatiser le traitement de documents. Le système devait rester simple à installer et suffisamment modulaire pour permettre des évolutions futures.",
            "Les objectifs fonctionnels étaient d'assurer l'inscription et la connexion des utilisateurs, l'import de fichiers PDF ou images, l'extraction OCR, l'affichage des informations détectées et la recherche dans les documents.",
            "Les objectifs techniques étaient de structurer le code, de stocker les données dans une base SQLite, de sécuriser les fichiers importés et de vérifier le fonctionnement de l'application avec des tests automatisés.",
        ],
        "screenshot": "Liste des objectifs ou cahier des charges",
    },
    {
        "title": "Analyse des besoins et fonctionnalités",
        "paragraphs": [
            "L'utilisateur doit pouvoir créer un compte, se connecter et accéder à un tableau de bord personnel. Depuis ce tableau de bord, il peut consulter les documents importés, les filtrer et accéder à leur détail.",
            "Le formulaire d'import accepte les extensions autorisées par la configuration. Le fichier est renommé de manière sécurisée, enregistré dans le dossier prévu, puis associé à un enregistrement dans la base de données.",
            "Après traitement, l'application affiche le statut OCR, le numéro de facture, le fournisseur, la date, le montant TTC et, lorsque cela est pertinent, des informations de maintenance ou une anomalie détectée.",
        ],
        "screenshot": "Diagramme des cas d'utilisation",
    },
    {
        "title": "Méthodologie et organisation du travail",
        "paragraphs": [
            "Le développement a été réalisé progressivement. Une première étape a consisté à comprendre le besoin et à définir les principales entités manipulées. L'architecture Flask et la structure des dossiers ont ensuite été mises en place.",
            "Les fonctionnalités ont été développées par modules : modèles de données, routes, traitement OCR, templates HTML et styles CSS. Chaque ajout a été vérifié localement afin de limiter les régressions.",
            "Cette organisation a permis de séparer la logique métier, l'accès aux données et l'interface utilisateur, ce qui facilite la maintenance et la compréhension du projet.",
        ],
        "screenshot": "Organisation des tâches ou arborescence du projet",
    },
    {
        "title": "Architecture technique",
        "paragraphs": [
            "Le backend repose sur Flask, un framework Python adapté au développement d'applications web. Les routes reçoivent les requêtes, appliquent les règles métier et renvoient les pages HTML ou les réponses JSON nécessaires.",
            "Flask-SQLAlchemy assure la communication avec la base SQLite. Les templates Jinja construisent les pages à partir des données récupérées par le serveur, tandis que Bootstrap et la feuille de styles personnalisée assurent la présentation responsive.",
            "Le traitement des images est réalisé avec Pillow et pytesseract. Les fichiers PDF sont convertis en images par pdf2image avant d'être transmis à Tesseract.",
        ],
        "screenshot": "Architecture technique de l'application",
    },
    {
        "title": "Conception de la base de données",
        "paragraphs": [
            "Le modèle Document constitue le coeur fonctionnel de l'application. Il conserve le nom et le chemin du fichier, le type de document, la date d'importation, le texte OCR et le statut du traitement.",
            "Les champs extraits comprennent le numéro de facture, le fournisseur, la date de facture, le montant TTC, le type de panne et la durée de maintenance. Le modèle contient également les informations liées aux anomalies et à la langue OCR.",
            "Les modèles User, ContactMessage et ComplianceAlert complètent le système en prenant respectivement en charge les comptes utilisateurs, les messages de contact et les alertes de conformité.",
        ],
        "screenshot": "Schéma relationnel ou modèles SQLAlchemy",
    },
    {
        "title": "Import et traitement des documents",
        "paragraphs": [
            "La route d'import vérifie la présence du fichier et contrôle son extension. Le nom est nettoyé avec secure_filename afin de réduire les risques liés aux noms de fichiers transmis par un utilisateur.",
            "Un document est d'abord créé avec le statut en_attente. Le moteur OCR extrait ensuite le texte. Les règles d'analyse identifient les champs métier et le document est classé comme facture, rapport de maintenance ou document technique.",
            "En cas de succès, le statut devient ok. Si une erreur survient ou si aucun contenu exploitable n'est extrait, le document est conservé avec le statut echec afin de permettre une vérification ultérieure.",
        ],
        "screenshot": "Formulaire d'import et message de résultat",
    },
    {
        "title": "OCR et extraction des informations",
        "paragraphs": [
            "Le module ocr_utils.py centralise les fonctions de reconnaissance et d'analyse. Pour une image, celle-ci est convertie en niveaux de gris avant le traitement. Pour un PDF, chaque page est convertie en image puis analysée.",
            "L'extraction des champs repose sur des expressions régulières. Elles recherchent notamment les motifs associés au numéro de facture, aux dates et au montant TTC. Des heuristiques permettent également d'identifier le fournisseur et le logo correspondant.",
            "Cette approche est adaptée à un prototype et reste facilement configurable. Sa précision dépend toutefois de la qualité du document et de la diversité des formats rencontrés.",
        ],
        "screenshot": "Exemple de texte OCR et données extraites",
    },
    {
        "title": "Tableau de bord et consultation",
        "paragraphs": [
            "Le tableau de bord présente les documents classés par date et fournit des indicateurs sur le nombre total de documents, les traitements réussis, les traitements en attente et les anomalies.",
            "Une recherche peut être effectuée sur le nom du fichier, le numéro de facture, le fournisseur ou le montant. Des filtres supplémentaires permettent de sélectionner le type de document, le statut OCR et une période.",
            "L'export CSV permet de réutiliser les informations dans un tableur ou un autre outil de gestion. La page de détail affiche les champs extraits et les informations brutes produites par l'OCR.",
        ],
        "screenshot": "Tableau de bord avec filtres",
    },
    {
        "title": "Détection des anomalies et assistant",
        "paragraphs": [
            "Le système peut signaler un montant inhabituel en le comparant à la moyenne des montants déjà enregistrés. Cette règle constitue une aide à la vérification et ne remplace pas une validation comptable.",
            "Les alertes de conformité sont associées au document concerné et peuvent être suivies depuis l'application. Elles permettent de rendre visibles les documents nécessitant une intervention.",
            "Un assistant est également disponible. Il fonctionne en mode local lorsque l'API externe n'est pas configurée et peut utiliser OpenAI lorsque la clé correspondante est disponible. Un retour local est prévu en cas d'erreur de l'API.",
        ],
        "screenshot": "Alerte de conformité et assistant virtuel",
    },
    {
        "title": "Sécurité et gestion des utilisateurs",
        "paragraphs": [
            "Les pages sensibles sont protégées par une vérification de session réalisée avant les requêtes. Les pages de connexion, d'inscription et certaines routes publiques restent accessibles sans session.",
            "Les mots de passe ne sont pas stockés en clair : ils sont transformés en empreinte à l'aide des fonctions de hachage de Werkzeug. Les noms de fichiers sont sécurisés avant leur enregistrement.",
            "Pour une mise en production, il serait nécessaire de compléter cette base avec une protection CSRF, une gestion plus fine des rôles, une configuration stricte des secrets et un stockage externe des fichiers.",
        ],
        "screenshot": "Page de connexion ou gestion des comptes",
    },
    {
        "title": "Interface utilisateur et expérience",
        "paragraphs": [
            "Les pages HTML héritent d'une structure commune définie dans base.html. Cette organisation garantit la cohérence de la navigation, des messages flash et du composant d'assistance.",
            "Les écrans principaux comprennent l'accueil, l'import, le détail d'un document, la FAQ, le contact et les pages d'administration. Le CSS personnalisé complète Bootstrap pour adapter l'affichage aux écrans de différentes tailles.",
            "Une attention particulière a été portée à la lisibilité des statuts et des messages d'erreur afin que l'utilisateur comprenne rapidement le résultat d'une action.",
        ],
        "screenshot": "Écrans principaux de l'application",
    },
    {
        "title": "Tests et validation",
        "paragraphs": [
            "Le projet contient des tests automatisés exécutables avec unittest. Ils vérifient notamment l'extraction des informations d'une facture, la redirection des utilisateurs non authentifiés et la création d'un document après import.",
            "Les tests couvrent aussi le fonctionnement de l'inscription, de la connexion, de la déconnexion, du mapping des logos fournisseurs et de l'assistant virtuel, y compris son retour en mode dégradé.",
            "La commande de validation est : python -m unittest tests.test_app -v. Des tests complémentaires sur plusieurs formats réels de factures seraient nécessaires avant un déploiement en production.",
        ],
        "screenshot": "Résultat d'exécution des tests",
    },
    {
        "title": "Difficultés rencontrées et solutions",
        "paragraphs": [
            "La première difficulté concerne la variabilité des documents. Une facture peut présenter des libellés, des polices ou des dispositions différentes. Des expressions régulières et des heuristiques ont été utilisées pour obtenir une première extraction générique.",
            "La configuration de Tesseract et de Poppler peut également varier selon le système d'exploitation. Le projet prévoit donc des paramètres de configuration pour indiquer le chemin du moteur OCR et documente les prérequis d'installation.",
            "Enfin, le traitement OCR peut échouer. L'application conserve alors le document et informe l'utilisateur au lieu d'interrompre complètement le workflow.",
        ],
        "screenshot": "Exemple d'erreur traitée par l'application",
    },
    {
        "title": "Compétences acquises et contribution personnelle",
        "paragraphs": [
            "Ce projet m'a permis de renforcer mes compétences en Python, Flask, SQLAlchemy, traitement de fichiers et développement d'interfaces web avec HTML, Jinja et CSS.",
            "J'ai également développé une meilleure compréhension de la conception d'une base de données, de la gestion des sessions, de la validation des entrées et de la mise en place de tests automatisés.",
            "Ma contribution a porté sur la conception et l'intégration des fonctionnalités de gestion documentaire, le traitement OCR, l'extraction des données et l'amélioration de l'expérience utilisateur.",
        ],
        "screenshot": "Fonctionnalité développée pendant le stage",
    },
    {
        "title": "Perspectives d'évolution",
        "paragraphs": [
            "Une version future pourrait intégrer une file de traitement asynchrone avec Celery ou RQ afin d'éviter de bloquer la requête pendant l'OCR de documents volumineux.",
            "La base SQLite pourrait être remplacée par PostgreSQL et les fichiers pourraient être stockés dans un service objet. Une pagination serait également utile lorsque le nombre de documents augmente.",
            "L'extraction pourrait être améliorée par des modèles spécialisés dans les factures, une validation humaine des champs et l'ajout de formats d'export supplémentaires. Une gestion plus complète des rôles et des journaux d'audit renforcerait aussi l'usage professionnel.",
        ],
        "screenshot": "Feuille de route des évolutions",
    },
    {
        "title": "Conclusion",
        "paragraphs": [
            "La réalisation de cette application a permis de répondre au besoin de centralisation et d'automatisation du traitement des factures et documents techniques. Le prototype obtenu propose un parcours complet, depuis l'import jusqu'à la consultation et à l'export des données.",
            "Le projet montre l'intérêt de combiner une application web, une base de données et l'OCR pour réduire les tâches répétitives. Il constitue une base fonctionnelle qui pourra être renforcée avant un déploiement à grande échelle.",
            "Ce stage a été une expérience formatrice, à la fois sur le plan technique et sur le plan de la conduite d'un projet logiciel, de l'analyse du besoin jusqu'à la validation des fonctionnalités.",
        ],
        "screenshot": "Capture finale de l'application",
    },
]


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)


def add_title_page(document):
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run("RAPPORT DE STAGE")
    run.bold = True
    run.font.size = Pt(28)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run("Développement d'une application web de gestion automatisée des factures")
    run.bold = True
    run.font.size = Pt(16)

    description = document.add_paragraph()
    description.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = description.add_run("Projet : Gestion Factures et Documents Techniques")
    run.italic = True
    run.font.size = Pt(13)

    for _ in range(4):
        document.add_paragraph()
    for label in [
        "Étudiant(e) : [À compléter]",
        "Formation : [À compléter]",
        "Organisme d'accueil : [À compléter]",
        "Encadrant(e) : [À compléter]",
        "Période du stage : [À compléter]",
        "Année universitaire : [À compléter]",
    ]:
        paragraph = document.add_paragraph(label)
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    document.add_page_break()


def add_toc(document):
    document.add_heading("Table des matières", level=1)
    for index, section in enumerate(sections, start=1):
        paragraph = document.add_paragraph(f"{index}. {section['title']}")
        paragraph.paragraph_format.left_indent = Inches(0.2)
    document.add_page_break()


def add_section(document, section):
    document.add_heading(section["title"], level=1)
    for text in section["paragraphs"]:
        document.add_paragraph(text)
    placeholder = document.add_paragraph(
        f"[Zone réservée pour capture d'écran : {section['screenshot']}]"
    )
    placeholder.runs[0].italic = True
    placeholder.runs[0].font.size = Pt(10)
    document.add_page_break()


def main():
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_toc(document)
    for section in sections:
        add_section(document, section)
    document.save(OUTPUT_PATH)
    print(f"Rapport généré : {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
