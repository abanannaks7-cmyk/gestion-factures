from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
import os

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUTPUT = os.path.join(BASE, "documentation_complete_projet_gestion_factures.docx")


def setup(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 24), ("Heading 1", 18), ("Heading 2", 13)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(28, 63, 91)


def title(document, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(28, 63, 91)


def heading(document, text, level=1):
    document.add_heading(text, level=level)


def paragraph(document, text):
    document.add_paragraph(text)


def bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def code(document, text):
    for line in text.strip("\n").splitlines():
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)


def page(document):
    document.add_page_break()


def build():
    document = Document()
    setup(document)

    title(document, "DOCUMENTATION COMPLÈTE DU PROJET")
    title(document, "Gestion automatisée des factures et documents techniques")
    paragraph(document, "Version : 1.0")
    paragraph(document, "Technologies principales : Python, Flask, SQLite, SQLAlchemy, Tesseract OCR")
    paragraph(document, "Ce document rassemble l'idée du projet, son fonctionnement, son code important, son installation, son utilisation et ses perspectives.")
    paragraph(document, "Auteur : [À compléter]")
    paragraph(document, "Date : [À compléter]")
    page(document)

    heading(document, "1. Présentation générale")
    paragraph(document, "Gestion Factures est une application web destinée à centraliser les factures et les documents techniques. Elle permet à un utilisateur de déposer un PDF ou une image, d'en extraire automatiquement le texte et de consulter les informations détectées dans une interface web.")
    heading(document, "1.1 Idée du projet", 2)
    paragraph(document, "L'idée est de remplacer une partie du classement et de la saisie manuels par un traitement numérique. Une facture importée est enregistrée, analysée par OCR, classée et rendue consultable avec ses champs importants.")
    heading(document, "1.2 Problème traité", 2)
    bullets(document, [
        "Les factures sont difficiles à retrouver lorsqu'elles sont dispersées dans plusieurs dossiers.",
        "La saisie manuelle du numéro, de la date et du montant prend du temps.",
        "Les documents techniques contiennent aussi des informations utiles qui peuvent être recherchées.",
        "Les erreurs de saisie et les montants inhabituels doivent être repérés plus facilement.",
    ])
    heading(document, "1.3 Solution proposée", 2)
    bullets(document, [
        "Une interface web avec authentification.",
        "Un import sécurisé des fichiers PDF, PNG, JPG et JPEG.",
        "Un OCR multilingue français, anglais et arabe.",
        "Une extraction automatique du numéro, fournisseur, date et montant TTC.",
        "Un tableau de bord avec recherche, filtres et export CSV.",
        "Une détection d'anomalies et un assistant de recherche.",
    ])

    heading(document, "2. Fonctionnalités du projet")
    features = [
        ("Authentification", "Création de compte, connexion, déconnexion et protection des pages privées."),
        ("Import", "Dépôt de fichiers avec contrôle des extensions et limite de taille à 16 Mo."),
        ("OCR", "Lecture de documents image et PDF avec Tesseract et pdf2image."),
        ("Extraction", "Recherche de motifs pour identifier les informations métier."),
        ("Classification", "Classement en facture, rapport de maintenance ou document technique."),
        ("Tableau de bord", "Statistiques, recherche par mot-clé, filtres par type, statut et dates."),
        ("Détail", "Affichage du fichier, des champs extraits et du texte OCR brut."),
        ("Export", "Téléchargement des documents et données principales au format CSV."),
        ("Anomalies", "Signalement des montants inhabituellement élevés et alertes de conformité."),
        ("Assistant", "Recherche locale et possibilité d'utiliser OpenAI si une clé est configurée."),
        ("Administration", "Gestion des alertes et des messages de contact."),
    ]
    for name, description in features:
        paragraph(document, f"{name} : {description}")

    heading(document, "3. Architecture du projet")
    paragraph(document, "L'application suit une architecture Flask classique. Le navigateur envoie une requête aux routes Flask. Les routes utilisent les modèles SQLAlchemy et les fonctions OCR, puis renvoient un template HTML ou une réponse JSON.")
    code(document, """
gestion-factures/
├── app/
│   ├── __init__.py          # création de l'application Flask
│   ├── config.py            # configuration et variables d'environnement
│   ├── models.py            # User, Document, ContactMessage, ComplianceAlert
│   ├── ocr_utils.py         # OCR, nettoyage et extraction des champs
│   ├── routes.py            # routes web, upload, export, assistant, admin
│   ├── templates/           # pages HTML Jinja
│   ├── static/css/          # styles CSS
│   ├── static/logos/        # logos fournisseurs
│   └── uploads/             # fichiers importés
├── scripts/                 # scripts de rapports et documentation
├── tests/test_app.py        # tests automatisés
├── requirements.txt         # dépendances Python
├── run.py                   # point d'entrée du serveur
└── database.db              # base SQLite créée au lancement
""")
    heading(document, "3.1 Flux principal", 2)
    code(document, """
Utilisateur -> /upload -> vérification du fichier
          -> sauvegarde sécurisée -> création du Document
          -> Tesseract OCR -> extraction des champs
          -> classification et anomalie -> base de données
          -> page de détail du document
""")

    heading(document, "4. Installation complète")
    heading(document, "4.1 Prérequis", 2)
    bullets(document, [
        "Python 3.11 ou version supérieure.",
        "Visual Studio Code avec les extensions Python et Pylance.",
        "Tesseract OCR avec les langues fra, eng et ara.",
        "Poppler pour convertir les fichiers PDF en images.",
    ])
    heading(document, "4.2 Windows", 2)
    code(document, r"""
cd C:\Users\Abdea\Downloads\gestion-factures\gestion-factures
python -m venv venv
venv\Scripts\activate
pip install -r requirements.txt
python run.py
""")
    paragraph(document, "Si Tesseract n'est pas dans le PATH Windows, créer un fichier .env à la racine du projet :")
    code(document, r"""
TESSERACT_CMD=C:\Program Files\Tesseract-OCR\tesseract.exe
SECRET_KEY=remplacer-par-une-cle-secrete
""")
    heading(document, "4.3 macOS et Linux", 2)
    code(document, """
python3 -m venv venv
source venv/bin/activate
pip install -r requirements.txt
python run.py
""")
    heading(document, "4.4 Accès au site", 2)
    paragraph(document, "Après le démarrage, ouvrir http://127.0.0.1:5000 dans le navigateur. Créer un compte, se connecter, puis importer un document test.")

    heading(document, "5. Configuration")
    paragraph(document, "La configuration se trouve dans app/config.py. Les paramètres sensibles doivent être placés dans un fichier .env et ne doivent pas être publiés dans Git.")
    code(document, """
SECRET_KEY = clé de session Flask
DATABASE_URL = URL de la base, SQLite par défaut
UPLOAD_FOLDER = dossier des fichiers importés
ALLOWED_EXTENSIONS = pdf, png, jpg, jpeg
MAX_CONTENT_LENGTH = 16 Mo
OCR_LANGUAGES = fra+eng+ara
TESSERACT_CMD = chemin vers tesseract.exe
OPENAI_API_KEY = clé facultative de l'assistant
OPENAI_MODEL = modèle OpenAI facultatif
ADMIN_USER / ADMIN_PASS = accès administration
SMTP_* = paramètres facultatifs d'e-mail
""")
    paragraph(document, "Attention : la valeur par défaut ADMIN_PASS=password et la SECRET_KEY de développement doivent être changées avant une mise en production.")

    heading(document, "6. Modèles de données")
    heading(document, "6.1 User", 2)
    paragraph(document, "Le modèle User contient l'identifiant, le nom d'utilisateur, l'e-mail, le mot de passe haché et la date de création. Les fonctions set_password et check_password gèrent le hachage et la vérification.")
    heading(document, "6.2 Document", 2)
    paragraph(document, "Document est l'entité centrale. Il contient le fichier original, le type, la date d'importation, le texte OCR, le statut et les champs extraits. Il contient aussi les informations de fournisseur, de maintenance, de langue et d'anomalie.")
    code(document, """
class Document(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom_fichier = db.Column(db.String(255), nullable=False)
    chemin_fichier = db.Column(db.String(500), nullable=False)
    type_document = db.Column(db.String(50), default="facture")
    texte_ocr = db.Column(db.Text, nullable=True)
    statut_ocr = db.Column(db.String(20), default="en_attente")
    numero_facture = db.Column(db.String(100), nullable=True)
    fournisseur = db.Column(db.String(255), nullable=True)
    date_facture = db.Column(db.String(50), nullable=True)
    montant_ttc = db.Column(db.String(50), nullable=True)
    anomalie = db.Column(db.Boolean, default=False)
""")
    heading(document, "6.3 Autres modèles", 2)
    paragraph(document, "ContactMessage stocke les messages envoyés par le formulaire de contact. ComplianceAlert stocke les règles, niveaux, messages et états des alertes associées à un document.")

    heading(document, "7. Fonctionnement OCR et extraction")
    paragraph(document, "Pour une image, Pillow convertit le document en niveaux de gris puis Tesseract réalise la reconnaissance. Pour un PDF, pdf2image convertit chaque page en image avant l'appel à Tesseract.")
    code(document, """
def extraire_texte(chemin_fichier, tesseract_cmd=None, langue="fra+eng+ara"):
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    extension = chemin_fichier.rsplit(".", 1)[-1].lower()
    if extension == "pdf":
        pages = convert_from_path(chemin_fichier, dpi=300)
        return "\\n".join(
            pytesseract.image_to_string(page, lang=langue) for page in pages
        ).strip()
    image = Image.open(chemin_fichier).convert("L")
    return pytesseract.image_to_string(image, lang=langue).strip()
""")
    paragraph(document, "Les expressions régulières recherchent le numéro après Facture ou Invoice, les dates au format jj/mm/aaaa, le montant après Total TTC ou Net à payer, puis une ligne probable pour le fournisseur.")
    code(document, """
texte = '''
SOCIETE ABC
Facture N° 2024-00123
Date : 15/02/2024
Total TTC : 1 250,00 MAD
'''
# Résultat attendu :
# numero_facture = 2024-00123
# date_facture = 15/02/2024
# montant_ttc = 1 250,00
# fournisseur = SOCIETE ABC
""")

    heading(document, "8. Routes et parcours utilisateur")
    bullets(document, [
        "/ : tableau de bord et filtres.",
        "/login : connexion utilisateur.",
        "/signup : création d'un compte.",
        "/logout : fermeture de session.",
        "/upload : import et traitement OCR.",
        "/export : export CSV des documents.",
        "/assistant/query : réponse JSON de l'assistant.",
        "Routes de détail, suppression, contact et administration dans app/routes.py.",
    ])
    heading(document, "8.1 Exemple du traitement d'un upload", 2)
    code(document, """
if fichier and extension_autorisee(fichier.filename):
    nom_securise = secure_filename(fichier.filename)
    chemin = os.path.join(UPLOAD_FOLDER, nom_securise)
    fichier.save(chemin)
    doc = Document(nom_fichier=nom_securise,
                   chemin_fichier=chemin,
                   statut_ocr="en_attente")
    db.session.add(doc)
    db.session.commit()
    texte = extraire_texte(chemin, langue=langue)
    champs = extraire_champs_facture(texte)
    doc.texte_ocr = texte
    doc.numero_facture = champs["numero_facture"]
    doc.fournisseur = champs["fournisseur"]
    doc.montant_ttc = champs["montant_ttc"]
    doc.statut_ocr = "ok" if texte else "echec"
    db.session.commit()
""")

    heading(document, "9. Sécurité et bonnes pratiques")
    bullets(document, [
        "Les pages privées vérifient la session avant de répondre.",
        "Les mots de passe sont hachés avec Werkzeug et ne sont pas stockés en clair.",
        "secure_filename nettoie le nom des fichiers importés.",
        "La taille maximale d'un upload est de 16 Mo.",
        "Les erreurs OCR sont capturées et affichées sans perdre le document.",
        "Les clés et mots de passe doivent être configurés dans .env.",
    ])
    paragraph(document, "Pour la production, ajouter Flask-Login, une protection CSRF, des rôles précis, une validation plus stricte des fichiers, des journaux d'audit et un serveur WSGI comme Gunicorn derrière Nginx.")

    heading(document, "10. Tests et validation")
    code(document, """
python -m unittest tests.test_app -v
""")
    paragraph(document, "Les tests actuels vérifient l'extraction OCR, la redirection des utilisateurs non authentifiés, la création d'un document, le mapping des logos, l'inscription, la connexion, la déconnexion et le fonctionnement de l'assistant avec retour local.")
    heading(document, "10.1 Exemple de test OCR", 2)
    code(document, """
def test_extraction_ocr_parse_facture(self):
    champs = extraire_champs_facture(texte)
    self.assertEqual(champs["numero_facture"], "2024-00123")
    self.assertEqual(champs["date_facture"], "15/02/2024")
    self.assertEqual(champs["fournisseur"], "SOCIETE ABC")
""")

    heading(document, "11. Dépannage")
    problems = [
        ("Tesseract introuvable", "Installer Tesseract et définir TESSERACT_CMD dans .env."),
        ("PDF non traité", "Installer Poppler et ajouter son dossier bin au PATH."),
        ("Port 5000 occupé", "Arrêter l'autre serveur ou modifier le port dans run.py."),
        ("OCR vide", "Utiliser une image plus nette, vérifier la langue et le contraste."),
        ("Erreur OpenAI", "L'assistant utilise automatiquement le mode local si l'API échoue."),
        ("Document verrouillé", "Fermer le fichier Word ou le programme qui l'utilise avant de le remplacer."),
    ]
    for problem, solution in problems:
        paragraph(document, f"{problem} : {solution}")

    heading(document, "12. Limites actuelles")
    bullets(document, [
        "Les règles OCR dépendent de la qualité et du format des factures.",
        "SQLite et le stockage local sont surtout adaptés au développement et à un usage limité.",
        "Le traitement OCR est synchrone et peut ralentir l'import d'un PDF volumineux.",
        "Les rôles utilisateurs et la protection CSRF doivent être renforcés en production.",
        "Le document importé doit être vérifié manuellement lorsque l'extraction est incertaine.",
    ])

    heading(document, "13. Évolutions recommandées")
    bullets(document, [
        "Ajouter une pagination et une recherche plus avancée.",
        "Utiliser Celery ou RQ pour traiter les documents en arrière-plan.",
        "Migrer vers PostgreSQL pour un environnement multi-utilisateur.",
        "Utiliser S3 ou Azure Blob pour les fichiers importés.",
        "Ajouter PDF.js pour la prévisualisation avancée.",
        "Ajouter une validation humaine champ par champ.",
        "Créer des règles de conformité configurables par l'administrateur.",
        "Ajouter des tests avec de vraies factures anonymisées.",
    ])

    heading(document, "14. Présentation orale du projet")
    paragraph(document, "Pour présenter le projet en quelques minutes, expliquer d'abord le problème de la gestion manuelle des factures. Montrer ensuite le parcours : connexion, import d'un fichier, OCR, affichage des champs, filtre dans le tableau de bord et export CSV.")
    paragraph(document, "La valeur principale du projet est le passage d'un document visuel à des informations structurées et consultables. Mentionner les technologies utilisées, les difficultés liées à l'OCR et les évolutions possibles.")
    heading(document, "14.1 Démonstration conseillée", 2)
    bullets(document, [
        "Ouvrir http://127.0.0.1:5000.",
        "Créer un utilisateur et se connecter.",
        "Importer une facture PDF ou JPG lisible.",
        "Montrer les champs extraits dans la page de détail.",
        "Retourner au tableau de bord et utiliser un filtre.",
        "Télécharger l'export CSV.",
        "Interroger l'assistant sur une facture ou une anomalie.",
    ])

    heading(document, "15. Références")
    bullets(document, [
        "https://flask.palletsprojects.com/",
        "https://flask-sqlalchemy.palletsprojects.com/",
        "https://tesseract-ocr.github.io/",
        "https://pypi.org/project/pytesseract/",
        "https://pypi.org/project/pdf2image/",
        "https://docs.python.org/3/",
    ])

    heading(document, "16. Checklist de livraison")
    bullets(document, [
        "Compléter les informations de l'auteur et de l'entreprise.",
        "Installer Python, Tesseract et Poppler.",
        "Créer l'environnement virtuel et installer requirements.txt.",
        "Configurer .env avec une SECRET_KEY réelle.",
        "Tester la connexion et l'import d'une facture.",
        "Exécuter python -m unittest tests.test_app -v.",
        "Remplacer les mots de passe de développement.",
        "Sauvegarder database.db et app/uploads selon la politique de l'entreprise.",
    ])

    document.save(OUTPUT)
    print(f"Documentation générée : {OUTPUT}")


if __name__ == "__main__":
    build()
